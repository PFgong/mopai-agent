import asyncio, io, json, os, tempfile, time, unittest, zipfile
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from common import unpack, pack
from engine import Config, convert, inventory, diagnostics
from compiler import compile_bytes

SAMPLE=Path(__file__).resolve().parents[2]/'dist/sample'

class AgentTests(unittest.TestCase):
    def test_zip_path_traversal_and_symlink_rejected(self):
        for name in ['../escape.tex','/tmp/escape.tex','folder\\escape.tex']:
            b=io.BytesIO()
            with zipfile.ZipFile(b,'w') as z:z.writestr(name,'test')
            with tempfile.TemporaryDirectory() as t:
                with self.assertRaises(ValueError):unpack(b.getvalue(),Path(t))
        b=io.BytesIO()
        with zipfile.ZipFile(b,'w') as z:
            info=zipfile.ZipInfo('link.tex');info.create_system=3;info.external_attr=0o120777<<16;z.writestr(info,'target')
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(ValueError):unpack(b.getvalue(),Path(t))

    def test_tex_project_nested_entry_and_resources(self):
        with tempfile.TemporaryDirectory() as t:
            root=Path(t); project=root/'project';project.mkdir()
            b=io.BytesIO()
            with zipfile.ZipFile(b,'w') as z:
                z.writestr('paper/main.tex',r'\documentclass{article}\begin{document}Test\end{document}')
                z.writestr('paper/refs.bib','@article{a,title={Title}}')
            f=root/'input.zip';f.write_bytes(b.getvalue())
            entry,warnings=convert(f,project,Config(entry='paper/main.tex'))
            self.assertEqual(entry,'paper/main.tex');self.assertTrue((project/'paper/refs.bib').exists())

    def test_docx_extraction_keeps_formula_and_image(self):
        from docx import Document
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls
        with tempfile.TemporaryDirectory() as t:
            root=Path(t); project=root/'out';project.mkdir();d=Document()
            d.add_heading('Conversion fixture',0); d.add_paragraph('Preserve this sentence.')
            p=d.add_paragraph();p._p.append(parse_xml('<m:oMath '+nsdecls('m')+'><m:r><m:t>x+1=2</m:t></m:r></m:oMath>'))
            d.add_picture(str(SAMPLE/'figures/attention.png'))
            f=root/'input.docx';d.save(f)
            entry,_=convert(f,project,Config(template='article'))
            tex=(project/entry).read_text()
            self.assertIn('Preserve this sentence',tex);self.assertIn('x+1=2',tex.replace(' ',''));self.assertIn('includegraphics',tex)
            self.assertTrue(any((project/'media').rglob('*.png')))

    def test_compiler_success_and_actual_error(self):
        good=compile_bytes(pack(SAMPLE,exclude=('latex-project.zip',)),'main.tex','xelatex')
        self.assertTrue(good['ok'],good['log'][-2000:]);self.assertIn('pdf_base64',good)
        with tempfile.TemporaryDirectory() as t:
            root=Path(t);(root/'main.tex').write_text(r'\documentclass{article}\begin{document}\unknowncommand\end{document}')
            bad=compile_bytes(pack(root),'main.tex','xelatex')
            self.assertFalse(bad['ok']);self.assertNotIn('pdf_base64',bad);self.assertIn('Undefined control sequence',bad['log'])

    def test_job_upload_compile_download_and_auth(self):
        import server
        with tempfile.TemporaryDirectory() as t,patch.object(server,'ROOT',Path(t)),patch.object(server,'API_TOKEN','test-secret'),patch.dict(os.environ,{'COMPILER_URL':'http://compiler'}):
            def compile_direct(project,entry,engine,url):
                import base64
                r=compile_bytes(pack(project),entry,engine)
                if r.get('pdf_base64'):(project/'output.pdf').write_bytes(base64.b64decode(r.pop('pdf_base64')))
                return r
            with patch.object(server,'compile_project',compile_direct),TestClient(server.app) as client:
                self.assertEqual(client.post('/jobs',files={'file':('input.tex',b'x')}).status_code,401)
                headers={'X-Agent-Token':'test-secret'}
                tex=r'\documentclass{article}\begin{document}Agent integration check.\end{document}'
                r=client.post('/jobs',headers=headers,files={'file':('input.tex',tex)},data={'config':json.dumps({'llm':{'api_key':'not-in-zip'}})})
                self.assertEqual(r.status_code,200,r.text);ident=r.json()['id']
                for _ in range(160):
                    job=client.get('/jobs/'+ident,headers=headers).json()
                    if job['status'] in {'completed','failed'}:break
                    time.sleep(.1)
                self.assertEqual(job['status'],'completed',job);self.assertEqual(job['report']['compile_status'],'passed');self.assertEqual(job['report']['content_status'],'needs_review')
                data=client.get('/jobs/'+ident+'/download/zip',headers=headers).content
                with zipfile.ZipFile(io.BytesIO(data)) as z:
                    self.assertIn('output.pdf',z.namelist())
                    self.assertFalse(any(b'not-in-zip' in z.read(n) for n in z.namelist()))
                self.assertEqual(client.delete('/jobs/'+ident,headers=headers).status_code,200)
                self.assertEqual(client.get('/jobs/'+ident,headers=headers).status_code,404)

    def test_structural_warnings(self):
        self.assertEqual(inventory(r'\includegraphics{a}\begin{equation}x\end{equation}\cite{a}')['figures'],1)
        self.assertTrue(diagnostics(r'\ref{missing}','Missing character:'))

if __name__=='__main__':unittest.main()
